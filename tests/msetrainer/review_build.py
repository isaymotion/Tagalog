import json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
d=json.load(open('rows.json'))
F=lambda **k: Font(name='Arial', **k)
HEAD_FILL=PatternFill('solid', fgColor='1C1F45'); INPUT_FILL=PatternFill('solid', fgColor='FFF7CC')
thin=Side(style='thin', color='D9DCEF'); BORDER=Border(bottom=thin)
wb=Workbook()
how=wb.active; how.title='How to use'
COLS=['Where','Topic','Speaker','Tagalog (draft)','English','Status','Suggested Tagalog','Notes']
WID=[14,30,18,60,52,11,48,30]
def sheet(name, rows):
    ws=wb.create_sheet(name)
    for i,(h,w) in enumerate(zip(COLS,WID),1):
        c=ws.cell(row=1,column=i,value=h); c.font=F(bold=True,color='FFFFFF'); c.fill=HEAD_FILL; c.alignment=Alignment(vertical='center',wrap_text=True)
        ws.column_dimensions[c.column_letter].width=w
    for r,row in enumerate(rows,2):
        for i,v in enumerate(row,1):
            c=ws.cell(row=r,column=i,value=v); c.font=F(size=10,bold=(i==4)); c.alignment=Alignment(vertical='top',wrap_text=True); c.border=BORDER
        for i in (6,7,8):
            c=ws.cell(row=r,column=i); c.fill=INPUT_FILL; c.font=F(size=10); c.alignment=Alignment(vertical='top',wrap_text=True); c.border=BORDER
    n=len(rows)+1
    dv=DataValidation(type='list', formula1='"OK,Fix,Unsure"', allow_blank=True, showErrorMessage=True, errorTitle='Status', error='Choose OK, Fix, or Unsure.')
    ws.add_data_validation(dv); dv.add(f'F2:F{n}')
    ws.conditional_formatting.add(f'F2:F{n}', CellIsRule(operator='equal', formula=['"OK"'], fill=PatternFill('solid', fgColor='DDF4EC')))
    ws.conditional_formatting.add(f'F2:F{n}', CellIsRule(operator='equal', formula=['"Fix"'], fill=PatternFill('solid', fgColor='FBE1E8')))
    ws.freeze_panes='D2'; ws.auto_filter.ref=f'A1:H{n}'
    ws.page_setup.orientation='landscape'; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0; ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.print_title_rows='1:1'
    return n
n1=sheet('MSE Trainer', d['mse']); n2=sheet('Usap Tayo', d['ut'])
# How to use
how.column_dimensions['A'].width=24; how.column_dimensions['B'].width=16; how.column_dimensions['C'].width=16; how.column_dimensions['D'].width=16; how.column_dimensions['E'].width=16
how['A1']='Tagalog review: new lines from October 2026'; how['A1'].font=F(bold=True,size=14)
lines=['Every line below is a draft written for practice cases and has not yet been checked by a native-speaking clinician.',
 'For each row, set Status (yellow column) to OK, Fix, or Unsure. If Fix, write your wording in Suggested Tagalog; add anything else in Notes.',
 'Edit only the yellow columns (Status, Suggested Tagalog, Notes). The Where column tells Claude which case or phrase to change.',
 'MSE Trainer sheet: the Tagalog cases and the Ask it in Tagalog questions (MSE Trainer, js/data-tl.js).',
 'Usap Tayo sheet: the five new discharge questions (respectful and everyday forms), patient and family answers in the five discharge practice cases, and the Ramon admission MSE case.',
 'When you are done, send the file back and Claude will apply the fixes to both apps; the audio will re-record automatically.']
for i,t in enumerate(lines,3):
    c=how.cell(row=i,column=1,value=t); c.font=F(size=11); how.merge_cells(start_row=i,start_column=1,end_row=i,end_column=8); c.alignment=Alignment(wrap_text=True,vertical='top'); how.row_dimensions[i].height=32
r=10
how.cell(row=r,column=1,value='Example row (how a fix looks)').font=F(bold=True)
for i,h in enumerate(COLS,1):
    c=how.cell(row=r+1,column=i,value=h); c.font=F(bold=True,color='FFFFFF'); c.fill=HEAD_FILL
ex=['tl-sp02','Poverty of speech','Clinician','Ikuwento po ninyo ang isang araw ninyo sa bahay.','Tell me about a day at home.','Fix','Ikuwento po ninyo kung ano ang ginagawa ninyo sa isang araw sa bahay.','More natural in Bataan']
for i,v in enumerate(ex,1):
    c=how.cell(row=r+2,column=i,value=v); c.font=F(size=10,italic=True); c.alignment=Alignment(wrap_text=True,vertical='top')
    if i>=6: c.fill=INPUT_FILL
how.row_dimensions[r+2].height=45
for col,w in zip('FGH',[10,40,22]): how.column_dimensions[col].width=w
how.column_dimensions['D'].width=40; how.column_dimensions['E'].width=30
r=15
how.cell(row=r,column=1,value='Progress').font=F(bold=True,size=12)
for i,h in enumerate(['Sheet','Lines','OK','Fix','Unsure','Not yet reviewed'],1):
    c=how.cell(row=r+1,column=i,value=h); c.font=F(bold=True,color='FFFFFF'); c.fill=HEAD_FILL
for j,(name,n) in enumerate([('MSE Trainer',n1),('Usap Tayo',n2)]):
    row=r+2+j; q=f"'{name}'"
    how.cell(row=row,column=1,value=name)
    how.cell(row=row,column=2,value=f'=COUNTA({q}!D2:D{n})')
    how.cell(row=row,column=3,value=f'=COUNTIF({q}!F2:F{n},"OK")')
    how.cell(row=row,column=4,value=f'=COUNTIF({q}!F2:F{n},"Fix")')
    how.cell(row=row,column=5,value=f'=COUNTIF({q}!F2:F{n},"Unsure")')
    how.cell(row=row,column=6,value=f'=B{row}-C{row}-D{row}-E{row}')
    for i in range(1,7): how.cell(row=row,column=i).font=F(size=11)
row=r+4; how.cell(row=row,column=1,value='Total').font=F(bold=True)
for i,col in enumerate('BCDEF',2):
    c=how.cell(row=row,column=i,value=f'=SUM({col}{r+2}:{col}{r+3})'); c.font=F(bold=True)
for row_ in how.iter_rows(min_row=1,max_row=25):
    for c in row_:
        if c.font.name!='Arial': c.font=F(size=c.font.sz or 11,bold=c.font.b,italic=c.font.i,color=c.font.color)
wb.save('tagalog-review-2026-10.xlsx')
print('saved', n1-1, n2-1)
